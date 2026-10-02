"""
DOCX Document Extractor.
Extracts text and paragraphs from Microsoft Word .docx documents using python-docx.
Extracts from both body paragraphs and embedded tables.
"""

import io
from typing import Tuple, List
import docx


def extract_text_from_docx(file_bytes: bytes) -> Tuple[bool, str, List[str], str]:
    """
    Extract text from DOCX file bytes.
    Returns: (success, raw_text, paragraphs, error_message)
    """
    if not file_bytes or len(file_bytes) == 0:
        return False, "", [], "Empty DOCX file provided."

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs: List[str] = []

        # Extract regular paragraphs
        for p in doc.paragraphs:
            cleaned = p.text.strip()
            if cleaned:
                paragraphs.append(cleaned)

        # Extract table cells content
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))

        raw_text = "\n\n".join(paragraphs).strip()
        if not raw_text:
            return False, "", [], "DOCX document contains no readable text."

        return True, raw_text, paragraphs, ""

    except Exception as e:
        return False, "", [], f"Failed to extract DOCX text: {str(e)}"
